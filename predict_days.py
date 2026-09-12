from pathlib import Path
import json

import numpy as np
import pandas as pd

from catboost import CatBoostRegressor

from openpyxl import load_workbook
from openpyxl.styles import (
    Alignment,
    Font,
    PatternFill,
)
from openpyxl.utils import get_column_letter

from predict import (
    enrich_station_features,
    load_model_info,
    predict_prices,
    prepare_input_data,
)

from predict_excel import (
    auto_fill_station_name,
)


PROJECT_ROOT = Path(__file__).resolve().parent

INPUT_DIR = (
    PROJECT_ROOT
    / "data"
    / "input"
)

INPUT_FILE = (
    INPUT_DIR
    / "days_prediction_input.xlsx"
)

OUTPUT_DIR = (
    PROJECT_ROOT
    / "output"
)

OUTPUT_FILE = (
    OUTPUT_DIR
    / "days_predictions.xlsx"
)

DAYS_MODEL_FILE = (
    PROJECT_ROOT
    / "models"
    / "days_model.cbm"
)

DAYS_METRICS_FILE = (
    OUTPUT_DIR
    / "days_model_metrics.json"
)

INPUT_SHEET = "予測入力"
DESCRIPTION_SHEET = "項目説明"
RESULT_SHEET = "予測結果"
MODEL_INFO_SHEET = "モデル情報"


INPUT_COLUMNS = [
    "property_id",

    "city",
    "district_name",

    "station_name",
    "station_line",
    "station_latitude",
    "station_longitude",

    "area_m2",
    "floor_plan",
    "building_age",
    "structure",
    "renovation",
    "use",
    "city_planning",

    "coverage_ratio",
    "floor_area_ratio",

    "listing_date",
    "asking_price",
]


REQUIRED_INPUT_COLUMNS = [
    "property_id",
    "city",
    "district_name",
    "area_m2",
    "floor_plan",
    "building_age",
    "listing_date",
    "asking_price",
]


TEXT_COLUMNS = [
    "property_id",
    "city",
    "district_name",
    "station_name",
    "station_line",
    "floor_plan",
    "structure",
    "renovation",
    "use",
    "city_planning",
]


NUMERIC_COLUMNS = [
    "station_latitude",
    "station_longitude",
    "area_m2",
    "building_age",
    "coverage_ratio",
    "floor_area_ratio",
    "asking_price",
]


SAMPLE_ROWS = [
    {
        "property_id": "P001",
        "city": "足立区",
        "district_name": "千住",

        "station_name": "北千住",
        "station_line": "",
        "station_latitude": "",
        "station_longitude": "",

        "area_m2": 65.2,
        "floor_plan": "3LDK",
        "building_age": 12,
        "structure": "RC",
        "renovation": "未改装",
        "use": "住宅",
        "city_planning": "商業地域",

        "coverage_ratio": 80,
        "floor_area_ratio": 400,

        "listing_date": "2026-09-01",
        "asking_price": 78_000_000,
    },
    {
        "property_id": "P002",
        "city": "世田谷区",
        "district_name": "三軒茶屋",

        "station_name": "三軒茶屋",
        "station_line": "",
        "station_latitude": "",
        "station_longitude": "",

        "area_m2": 55.0,
        "floor_plan": "2LDK",
        "building_age": 8,
        "structure": "RC",
        "renovation": "未改装",
        "use": "住宅",
        "city_planning": "近隣商業地域",

        "coverage_ratio": 80,
        "floor_area_ratio": 300,

        "listing_date": "2026-09-01",
        "asking_price": 110_000_000,
    },
]


FIELD_DESCRIPTIONS = [
    (
        "property_id",
        "必須",
        "物件識別ID",
    ),
    (
        "city",
        "必須",
        "東京都の市区町村",
    ),
    (
        "district_name",
        "必須",
        "町・地区名",
    ),
    (
        "station_name",
        "推奨",
        "最寄駅。空欄の場合は地区情報から代表駅を補完",
    ),
    (
        "station_line",
        "任意",
        "路線。空欄の場合は駅名から補完",
    ),
    (
        "station_latitude",
        "任意",
        "駅緯度。空欄の場合は駅名から補完",
    ),
    (
        "station_longitude",
        "任意",
        "駅経度。空欄の場合は駅名から補完",
    ),
    (
        "area_m2",
        "必須",
        "専有面積（㎡）",
    ),
    (
        "floor_plan",
        "必須",
        "間取り",
    ),
    (
        "building_age",
        "必須",
        "売出時点の築年数",
    ),
    (
        "structure",
        "推奨",
        "建物構造",
    ),
    (
        "renovation",
        "任意",
        "改装状況",
    ),
    (
        "use",
        "任意",
        "物件用途",
    ),
    (
        "city_planning",
        "推奨",
        "都市計画・用途地域",
    ),
    (
        "coverage_ratio",
        "任意",
        "建ぺい率",
    ),
    (
        "floor_area_ratio",
        "任意",
        "容積率",
    ),
    (
        "listing_date",
        "必須",
        "販売開始日",
    ),
    (
        "asking_price",
        "必須",
        "売出価格（円）",
    ),
]


def print_header(
    text: str,
) -> None:
    print()
    print("=" * 60)
    print(text)
    print("=" * 60)
    print()


def create_input_template() -> None:
    INPUT_DIR.mkdir(
        parents=True,
        exist_ok=True,
    )

    df = pd.DataFrame(
        SAMPLE_ROWS,
        columns=INPUT_COLUMNS,
    )

    description_df = pd.DataFrame(
        FIELD_DESCRIPTIONS,
        columns=[
            "項目名",
            "入力区分",
            "説明",
        ],
    )

    with pd.ExcelWriter(
        INPUT_FILE,
        engine="openpyxl",
    ) as writer:
        df.to_excel(
            writer,
            index=False,
            sheet_name=INPUT_SHEET,
        )

        description_df.to_excel(
            writer,
            index=False,
            sheet_name=DESCRIPTION_SHEET,
        )

    style_input_excel()

    print_header(
        "成約日数予測用Excelを作成しました"
    )

    print(
        f"保存先: "
        f"{INPUT_FILE}"
    )

    print()
    print(
        "内容を確認・編集してから"
        "再度実行してください。"
    )


def style_input_excel() -> None:
    workbook = load_workbook(
        INPUT_FILE
    )

    header_fill = PatternFill(
        fill_type="solid",
        fgColor="1F4E78",
    )

    header_font = Font(
        color="FFFFFF",
        bold=True,
    )

    required_fill = PatternFill(
        fill_type="solid",
        fgColor="FFF2CC",
    )

    recommended_fill = PatternFill(
        fill_type="solid",
        fgColor="E2F0D9",
    )

    for sheet_name in workbook.sheetnames:
        worksheet = workbook[
            sheet_name
        ]

        worksheet.freeze_panes = "A2"

        for cell in worksheet[1]:
            cell.fill = header_fill
            cell.font = header_font

            cell.alignment = Alignment(
                horizontal="center",
                vertical="center",
            )

        for column_cells in worksheet.columns:
            max_length = 0

            column_letter = (
                get_column_letter(
                    column_cells[0].column
                )
            )

            for cell in column_cells:
                if cell.value is None:
                    continue

                max_length = max(
                    max_length,
                    len(
                        str(
                            cell.value
                        )
                    ),
                )

            worksheet.column_dimensions[
                column_letter
            ].width = min(
                max(
                    max_length + 2,
                    12,
                ),
                45,
            )

    input_sheet = workbook[
        INPUT_SHEET
    ]

    header_map = {
        cell.value: cell.column
        for cell in input_sheet[1]
    }

    for column in REQUIRED_INPUT_COLUMNS:
        if column in header_map:
            input_sheet.cell(
                row=1,
                column=header_map[
                    column
                ],
            ).fill = required_fill

    recommended_columns = [
        "station_name",
        "structure",
        "city_planning",
    ]

    for column in recommended_columns:
        if column in header_map:
            input_sheet.cell(
                row=1,
                column=header_map[
                    column
                ],
            ).fill = recommended_fill

    workbook.save(
        INPUT_FILE
    )


def load_input_data():
    if not INPUT_FILE.exists():
        create_input_template()
        return None

    try:
        df = pd.read_excel(
            INPUT_FILE,
            sheet_name=INPUT_SHEET,
        )

    except ValueError as error:
        raise RuntimeError(
            f"'{INPUT_SHEET}' シートが"
            "見つかりません。"
        ) from error

    df = (
        df
        .dropna(
            how="all"
        )
        .reset_index(
            drop=True
        )
    )

    if df.empty:
        raise RuntimeError(
            "予測対象物件がありません。"
        )

    print(
        f"成約日数予測対象: "
        f"{len(df):,}件"
    )

    return df


def validate_input_columns(
    df: pd.DataFrame,
) -> None:
    missing = [
        column
        for column in REQUIRED_INPUT_COLUMNS
        if column not in df.columns
    ]

    if missing:
        raise KeyError(
            "必要な入力列が不足しています:\n- "
            + "\n- ".join(
                missing
            )
        )


def add_optional_columns(
    df: pd.DataFrame,
) -> pd.DataFrame:
    df = df.copy()

    for column in INPUT_COLUMNS:
        if column in df.columns:
            continue

        if column in TEXT_COLUMNS:
            df[column] = pd.Series(
                pd.NA,
                index=df.index,
                dtype="string",
            )

        else:
            df[column] = np.nan

    return df


def clean_input_data(
    df: pd.DataFrame,
) -> pd.DataFrame:
    df = df.copy()

    validate_input_columns(
        df
    )

    df = add_optional_columns(
        df
    )

    for column in TEXT_COLUMNS:
        df[column] = (
            df[column]
            .astype("string")
            .str.strip()
            .replace(
                {
                    "": pd.NA,
                    "nan": pd.NA,
                    "None": pd.NA,
                    "<NA>": pd.NA,
                }
            )
        )

    for column in NUMERIC_COLUMNS:
        df[column] = (
            pd.to_numeric(
                df[column],
                errors="coerce",
            )
        )

    df[
        "listing_date"
    ] = pd.to_datetime(
        df[
            "listing_date"
        ],
        errors="coerce",
    )

    errors = []

    for index, row in df.iterrows():
        excel_row = index + 2

        for column in [
            "property_id",
            "city",
            "district_name",
            "floor_plan",
        ]:
            if pd.isna(
                row[column]
            ):
                errors.append(
                    f"{excel_row}行目: "
                    f"{column} が空です。"
                )

        if (
            pd.isna(
                row["area_m2"]
            )
            or row["area_m2"] <= 0
        ):
            errors.append(
                f"{excel_row}行目: "
                "area_m2 が不正です。"
            )

        if (
            pd.isna(
                row["building_age"]
            )
            or row["building_age"] < 0
        ):
            errors.append(
                f"{excel_row}行目: "
                "building_age が不正です。"
            )

        if (
            pd.isna(
                row["asking_price"]
            )
            or row["asking_price"] <= 0
        ):
            errors.append(
                f"{excel_row}行目: "
                "asking_price が不正です。"
            )

        if pd.isna(
            row[
                "listing_date"
            ]
        ):
            errors.append(
                f"{excel_row}行目: "
                "listing_date が不正です。"
            )

    duplicate_mask = (
        df[
            "property_id"
        ].notna()
        & df[
            "property_id"
        ].duplicated(
            keep=False
        )
    )

    if duplicate_mask.any():
        duplicates = (
            df.loc[
                duplicate_mask,
                "property_id",
            ]
            .astype(str)
            .unique()
            .tolist()
        )

        errors.append(
            "property_id が重複しています: "
            + ", ".join(
                duplicates
            )
        )

    if errors:
        print()
        print(
            "入力エラー"
        )

        for error in errors[:20]:
            print(
                f"- {error}"
            )

        raise ValueError(
            "入力Excelを修正してください。"
        )

    return df


def create_date_features(
    df: pd.DataFrame,
) -> pd.DataFrame:
    df = df.copy()

    df[
        "listing_year"
    ] = (
        df[
            "listing_date"
        ]
        .dt.year
        .astype(int)
    )

    df[
        "listing_month"
    ] = (
        df[
            "listing_date"
        ]
        .dt.month
        .astype(int)
    )

    df[
        "listing_quarter"
    ] = (
        df[
            "listing_date"
        ]
        .dt.quarter
        .astype(int)
    )

    df[
        "asking_price_per_m2"
    ] = (
        df[
            "asking_price"
        ]
        / df[
            "area_m2"
        ]
    )

    return df


def create_price_features(
    df: pd.DataFrame,
) -> pd.DataFrame:
    df = df.copy()

    (
        price_model,
        price_model_info,
        price_features,
        price_categories,
    ) = load_model_info()

    print()
    print(
        "Ver.5価格AIによる"
        "適正価格を算出します。"
    )

    df = auto_fill_station_name(
        df
    )

    df = enrich_station_features(
        df,
        price_features,
    )

    df[
        "transaction_year"
    ] = df[
        "listing_year"
    ]

    df[
        "transaction_quarter"
    ] = df[
        "listing_quarter"
    ]

    price_input = (
        prepare_input_data(
            df,
            price_features,
            price_categories,
        )
    )

    (
        predicted_unit_price,
        predicted_price,
    ) = predict_prices(
        price_model,
        price_input,
        price_features,
    )

    df[
        "ai_price_per_m2"
    ] = predicted_unit_price

    df[
        "ai_estimated_price"
    ] = predicted_price

    df[
        "price_gap_amount"
    ] = (
        df[
            "asking_price"
        ]
        - df[
            "ai_estimated_price"
        ]
    )

    df[
        "price_gap_ratio"
    ] = (
        df[
            "price_gap_amount"
        ]
        / df[
            "ai_estimated_price"
        ]
    )

    df[
        "price_gap_ratio_percent"
    ] = (
        df[
            "price_gap_ratio"
        ]
        * 100
    )

    df[
        "price_model_version"
    ] = str(
        price_model_info.get(
            "version",
            "unknown",
        )
    )

    return df


def load_days_model():
    if not DAYS_MODEL_FILE.exists():
        print_header(
            "成約日数AI: 教師データ待ち"
        )

        print(
            "成約日数モデルは"
            "まだ学習されていません。"
        )

        print()
        print(
            "実成約履歴を100件以上入力後、"
        )

        print(
            "python main.py days-full"
        )

        print(
            "を実行してください。"
        )

        return None

    if not DAYS_METRICS_FILE.exists():
        raise FileNotFoundError(
            "days_model_metrics.json が"
            "見つかりません。"
        )

    with open(
        DAYS_METRICS_FILE,
        "r",
        encoding="utf-8",
    ) as file:
        model_info = json.load(
            file
        )

    feature_columns = (
        model_info.get(
            "features",
            [],
        )
    )

    categorical_columns = (
        model_info.get(
            "categorical_features",
            [],
        )
    )

    if not feature_columns:
        raise RuntimeError(
            "days_model_metrics.json に"
            "features がありません。"
        )

    model = CatBoostRegressor()

    model.load_model(
        str(
            DAYS_MODEL_FILE
        )
    )

    print()
    print(
        "成約日数AI読み込み完了"
    )

    print(
        f"Version: "
        f"{model_info.get('version', 'unknown')}"
    )

    test_metrics = (
        model_info.get(
            "test_metrics",
            {}
        )
    )

    if (
        "mae_days"
        in test_metrics
    ):
        print(
            f"Final Test MAE: "
            f"{test_metrics['mae_days']:.2f}日"
        )

    if (
        "mape_percent"
        in test_metrics
    ):
        print(
            f"Final Test MAPE: "
            f"{test_metrics['mape_percent']:.2f}%"
        )

    return (
        model,
        model_info,
        feature_columns,
        categorical_columns,
    )


def prepare_days_input(
    df: pd.DataFrame,
    feature_columns,
    categorical_columns,
) -> pd.DataFrame:
    df = df.copy()

    missing_columns = [
        column
        for column in feature_columns
        if column not in df.columns
    ]

    if missing_columns:
        raise KeyError(
            "成約日数AIに必要な特徴量が"
            "不足しています:\n- "
            + "\n- ".join(
                missing_columns
            )
        )

    for column in feature_columns:
        if (
            column
            in categorical_columns
        ):
            df[column] = (
                df[column]
                .fillna("不明")
                .astype(str)
            )

        else:
            df[column] = (
                pd.to_numeric(
                    df[column],
                    errors="coerce",
                )
            )

    return df


def predict_contract_days(
    model,
    df: pd.DataFrame,
    feature_columns,
):
    predicted_log = (
        model.predict(
            df[
                feature_columns
            ]
        )
    )

    predicted_days = (
        np.expm1(
            predicted_log
        )
    )

    predicted_days = (
        np.maximum(
            predicted_days,
            1,
        )
    )

    return predicted_days


def classify_days(
    days: int,
) -> str:
    if days <= 30:
        return "30日以内"

    if days <= 60:
        return "31〜60日"

    if days <= 90:
        return "61〜90日"

    return "91日以上"


def create_result(
    df: pd.DataFrame,
    predicted_days,
) -> pd.DataFrame:
    result = df.copy()

    result[
        "predicted_days_to_contract"
    ] = (
        np.round(
            predicted_days
        )
        .astype(int)
    )

    result[
        "predicted_contract_date"
    ] = (
        result[
            "listing_date"
        ]
        + pd.to_timedelta(
            result[
                "predicted_days_to_contract"
            ],
            unit="D",
        )
    )

    result[
        "predicted_contract_period"
    ] = (
        result[
            "predicted_days_to_contract"
        ]
        .apply(
            classify_days
        )
    )

    result[
        "ai_estimated_price"
    ] = (
        result[
            "ai_estimated_price"
        ]
        .round()
        .astype(int)
    )

    result[
        "ai_price_per_m2"
    ] = (
        result[
            "ai_price_per_m2"
        ]
        .round()
        .astype(int)
    )

    result[
        "price_gap_amount"
    ] = (
        result[
            "price_gap_amount"
        ]
        .round()
        .astype(int)
    )

    result[
        "price_gap_ratio_percent"
    ] = (
        result[
            "price_gap_ratio_percent"
        ]
        .round(2)
    )

    return result


def save_result(
    result: pd.DataFrame,
    days_model_info,
) -> None:
    OUTPUT_DIR.mkdir(
        parents=True,
        exist_ok=True,
    )

    model_info_rows = [
        {
            "項目": "成約日数モデル",
            "値": (
                days_model_info.get(
                    "model",
                    "CatBoostRegressor",
                )
            ),
        },
        {
            "項目": "Version",
            "値": (
                days_model_info.get(
                    "version",
                    "unknown",
                )
            ),
        },
        {
            "項目": "学習方式",
            "値": (
                "時間順 70% / 15% / 15%"
            ),
        },
        {
            "項目": "価格AI連携",
            "値": "あり",
        },
    ]

    test_metrics = (
        days_model_info.get(
            "test_metrics",
            {}
        )
    )

    if "mae_days" in test_metrics:
        model_info_rows.append(
            {
                "項目": "Final Test MAE",
                "値": (
                    f"{test_metrics['mae_days']:.2f}日"
                ),
            }
        )

    if "mape_percent" in test_metrics:
        model_info_rows.append(
            {
                "項目": "Final Test MAPE",
                "値": (
                    f"{test_metrics['mape_percent']:.2f}%"
                ),
            }
        )

    model_info_df = pd.DataFrame(
        model_info_rows
    )

    with pd.ExcelWriter(
        OUTPUT_FILE,
        engine="openpyxl",
    ) as writer:
        result.to_excel(
            writer,
            index=False,
            sheet_name=RESULT_SHEET,
        )

        model_info_df.to_excel(
            writer,
            index=False,
            sheet_name=MODEL_INFO_SHEET,
        )

    style_output_excel()

    print_header(
        "成約日数予測完了"
    )

    print(
        f"件数: "
        f"{len(result):,}件"
    )

    print(
        f"保存先: "
        f"{OUTPUT_FILE}"
    )


def style_output_excel() -> None:
    workbook = load_workbook(
        OUTPUT_FILE
    )

    header_fill = PatternFill(
        fill_type="solid",
        fgColor="1F4E78",
    )

    header_font = Font(
        color="FFFFFF",
        bold=True,
    )

    for sheet_name in workbook.sheetnames:
        worksheet = workbook[
            sheet_name
        ]

        worksheet.freeze_panes = "A2"

        for cell in worksheet[1]:
            cell.fill = header_fill
            cell.font = header_font

            cell.alignment = Alignment(
                horizontal="center",
                vertical="center",
            )

        for column_cells in worksheet.columns:
            max_length = 0

            column_letter = (
                get_column_letter(
                    column_cells[0].column
                )
            )

            for cell in column_cells:
                if cell.value is None:
                    continue

                max_length = max(
                    max_length,
                    len(
                        str(
                            cell.value
                        )
                    ),
                )

            worksheet.column_dimensions[
                column_letter
            ].width = min(
                max(
                    max_length + 2,
                    12,
                ),
                35,
            )

    result_sheet = workbook[
        RESULT_SHEET
    ]

    header_map = {
        cell.value: cell.column
        for cell in result_sheet[1]
    }

    date_columns = [
        "listing_date",
        "predicted_contract_date",
    ]

    for column in date_columns:
        column_index = (
            header_map.get(
                column
            )
        )

        if column_index is None:
            continue

        for row in range(
            2,
            result_sheet.max_row + 1,
        ):
            result_sheet.cell(
                row=row,
                column=column_index,
            ).number_format = (
                "yyyy-mm-dd"
            )

    price_columns = [
        "asking_price",
        "ai_estimated_price",
        "price_gap_amount",
    ]

    for column in price_columns:
        column_index = (
            header_map.get(
                column
            )
        )

        if column_index is None:
            continue

        for row in range(
            2,
            result_sheet.max_row + 1,
        ):
            result_sheet.cell(
                row=row,
                column=column_index,
            ).number_format = (
                '#,##0"円"'
            )

    workbook.save(
        OUTPUT_FILE
    )


def show_predictions(
    result: pd.DataFrame,
) -> None:
    print()
    print(
        "成約日数AI予測結果"
    )

    for index, row in (
        result.iterrows()
    ):
        print()
        print(
            "-" * 50
        )

        print(
            f"物件 {index + 1}"
        )

        print(
            f"ID: "
            f"{row['property_id']}"
        )

        print(
            f"地域: "
            f"{row['city']} "
            f"{row['district_name']}"
        )

        print(
            f"最寄駅: "
            f"{row['station_name']}"
        )

        print(
            f"売出価格: "
            f"{row['asking_price']:,.0f}円"
        )

        print(
            f"AI想定価格: "
            f"{row['ai_estimated_price']:,.0f}円"
        )

        print(
            f"価格乖離率: "
            f"{row['price_gap_ratio_percent']:+.2f}%"
        )

        print(
            f"予測成約日数: "
            f"{row['predicted_days_to_contract']}日"
        )

        print(
            f"予測成約日: "
            f"{row['predicted_contract_date'].date()}"
        )

        print(
            f"期間区分: "
            f"{row['predicted_contract_period']}"
        )


def main() -> None:
    print_header(
        "東京都中古マンション "
        "成約日数AI予測"
    )

    df = load_input_data()

    if df is None:
        return

    days_model_data = (
        load_days_model()
    )

    if days_model_data is None:
        return

    (
        days_model,
        days_model_info,
        days_features,
        days_categories,
    ) = days_model_data

    df = clean_input_data(
        df
    )

    df = create_date_features(
        df
    )

    df = create_price_features(
        df
    )

    days_input = (
        prepare_days_input(
            df,
            days_features,
            days_categories,
        )
    )

    predicted_days = (
        predict_contract_days(
            days_model,
            days_input,
            days_features,
        )
    )

    result = create_result(
        df,
        predicted_days,
    )

    save_result(
        result,
        days_model_info,
    )

    show_predictions(
        result
    )


if __name__ == "__main__":
    main()